"""Génération d'un DOCX texte simple depuis un fichier texte.

Usage:
    python text_to_docx.py <input.txt> <output.docx> [--title "Titre"]
"""

import argparse
import sys


def text_to_docx(input_txt: str, output_docx: str, title: str = "") -> None:
    from docx import Document

    with open(input_txt, "r", encoding="utf-8") as f:
        content = f.read()

    doc = Document()
    first = title or (content.splitlines()[0].strip() if content.splitlines() else "Document")
    doc.add_heading(first, level=1)

    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if line == content.splitlines()[0] and not title:
            continue
        doc.add_paragraph(stripped)

    doc.save(output_docx)
    print(f"DOCX généré : {output_docx}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Convertir un texte en DOCX simple.")
    parser.add_argument("input_txt", help="Fichier texte source")
    parser.add_argument("output_docx", help="Fichier DOCX de sortie")
    parser.add_argument("--title", default="", help="Titre du document")
    args = parser.parse_args()

    try:
        text_to_docx(args.input_txt, args.output_docx, args.title)
    except Exception as e:
        print(f"Erreur conversion texte vers DOCX : {e}", file=sys.stderr)
        sys.exit(1)